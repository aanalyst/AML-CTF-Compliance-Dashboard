"""
AML/CTF Compliance Dashboard Project
Step 4: AUSTRAC-style SAR/SMR Excel Report
Input:  amlnet.db
Output: AML_SAR_Report.xlsx
"""

import sqlite3
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

# ─────────────────────────────────────────────
# LOAD DATA FROM SQLITE
# ─────────────────────────────────────────────
conn = sqlite3.connect("amlnet.db")

confirmed   = pd.read_sql("SELECT * FROM vw_suspicious_transactions ORDER BY amount DESC", conn)
ml_flagged  = pd.read_sql("SELECT * FROM vw_ml_flagged ORDER BY ml_risk_score DESC LIMIT 100", conn)
ttr         = pd.read_sql("SELECT * FROM vw_ttr_breaches ORDER BY amount DESC", conn)
structuring = pd.read_sql("SELECT * FROM vw_structuring_candidates ORDER BY amount DESC", conn)
typology    = pd.read_sql("SELECT * FROM vw_typology_summary", conn)
conn.close()

# ─────────────────────────────────────────────
# STYLES
# ─────────────────────────────────────────────
HEADER_FILL    = PatternFill("solid", start_color="1F4E79")   # dark blue
SUBHEAD_FILL   = PatternFill("solid", start_color="2E75B6")   # mid blue
ALT_FILL       = PatternFill("solid", start_color="D6E4F0")   # light blue
WARNING_FILL   = PatternFill("solid", start_color="FFE699")   # amber
DANGER_FILL    = PatternFill("solid", start_color="FF9999")   # red

HEADER_FONT    = Font(name="Arial", bold=True, color="FFFFFF", size=11)
SUBHEAD_FONT   = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT      = Font(name="Arial", size=10)
TITLE_FONT     = Font(name="Arial", bold=True, size=14, color="1F4E79")
LABEL_FONT     = Font(name="Arial", bold=True, size=10, color="1F4E79")

CENTER  = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT    = Alignment(horizontal="left",   vertical="center", wrap_text=False)
RIGHT   = Alignment(horizontal="right",  vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

AUD = '#,##0.00'
PCT = '0.0000%'

def style_header_row(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill   = HEADER_FILL
        cell.font   = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER

def style_data_row(ws, row, cols, alt=False):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        if alt:
            cell.fill = ALT_FILL
        cell.font      = BODY_FONT
        cell.alignment = LEFT
        cell.border    = THIN_BORDER

def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def write_dataframe(ws, df, start_row, col_headers, col_widths, amount_cols=None, pct_cols=None):
    # Header
    for col, header in enumerate(col_headers, 1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.fill      = SUBHEAD_FILL
        cell.font      = SUBHEAD_FONT
        cell.alignment = CENTER
        cell.border    = THIN_BORDER

    # Data rows
    for r_idx, (_, row_data) in enumerate(df.iterrows(), start_row + 1):
        alt = (r_idx % 2 == 0)
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font      = BODY_FONT
            cell.alignment = LEFT
            cell.border    = THIN_BORDER
            if alt:
                cell.fill = ALT_FILL
            if amount_cols and c_idx in amount_cols:
                cell.number_format = AUD
                cell.alignment     = RIGHT
            if pct_cols and c_idx in pct_cols:
                cell.number_format = PCT
                cell.alignment     = RIGHT

    set_col_widths(ws, col_widths)
    return r_idx  # last row written

wb = Workbook()

# ─────────────────────────────────────────────
# SHEET 1: SUMMARY
# ─────────────────────────────────────────────
ws1 = wb.active
ws1.title = "Summary"
ws1.sheet_view.showGridLines = False
ws1.row_dimensions[1].height = 40
ws1.row_dimensions[2].height = 20

# Title block
ws1.merge_cells("A1:F1")
title_cell = ws1["A1"]
title_cell.value     = "AUSTRAC Suspicious Matter Report (SMR) — Summary"
title_cell.font      = TITLE_FONT
title_cell.alignment = CENTER
title_cell.fill      = PatternFill("solid", start_color="EBF3FB")

ws1.merge_cells("A2:F2")
sub_cell = ws1["A2"]
sub_cell.value     = f"Generated: {datetime.now().strftime('%d %B %Y')}  |  Dataset: AMLNet v2 (Synthetic Australian Transactions)  |  Reporting Period: Feb 2025 – Aug 2025"
sub_cell.font      = Font(name="Arial", size=9, italic=True, color="595959")
sub_cell.alignment = CENTER

ws1.append([])

# KPI section
kpi_headers = ["Metric", "Value", "Notes"]
kpi_data = [
    ["Total Transactions Reviewed",     f"{len(confirmed) + 1_088_427:,}",    "Full AMLNet v2 dataset"],
    ["Confirmed Laundering Cases",       f"{len(confirmed):,}",                "isMoneyLaundering = 1"],
    ["Total Laundering Value (AUD)",     f"${confirmed['amount'].sum():,.2f}",  "Sum of confirmed case amounts"],
    ["Average Laundering Amount (AUD)",  f"${confirmed['amount'].mean():,.2f}", "Mean transaction value"],
    ["Laundering Rate",                  "0.16%",                              "Confirmed cases / total transactions"],
    ["ML Flagged Transactions",          "21,938",                             "ml_risk_score >= 0.50"],
    ["TTR Breaches (≥$10,000 AUD)",     f"{len(ttr):,}",                      "AUSTRAC reporting obligation"],
    ["Structuring Candidates",           f"{len(structuring):,}",              "$9,000–$9,999 AUD band"],
    ["Dominant Laundering Typology",     "Layering",                           "1,370 of 1,745 cases"],
    ["Highest Risk City",                "Sydney",                             "By transaction volume"],
]

ws1.append([""] + kpi_headers)
style_header_row(ws1, ws1.max_row, 4)

for i, row in enumerate(kpi_data):
    ws1.append([""] + row)
    r = ws1.max_row
    alt = (i % 2 == 0)
    style_data_row(ws1, r, 4, alt)
    ws1.cell(r, 2).font = LABEL_FONT

ws1.append([])

# Typology breakdown title
ws1.append(["", "Laundering Typology Breakdown", "", "", "", ""])
style_header_row(ws1, ws1.max_row, 6)

# Typology column headers
ws1.append(["", "Typology", "Case Count", "Avg Amount (AUD)", "Total Amount (AUD)", ""])
style_header_row(ws1, ws1.max_row, 6)

typo_display = typology[typology["laundering_typology"] != "normal"][
    ["laundering_typology", "laundering_count", "avg_amount", "total_amount"]
].copy()
typo_display.columns = ["Typology", "Case Count", "Avg Amount (AUD)", "Total Amount (AUD)"]

for i, (_, row) in enumerate(typo_display.iterrows()):
    ws1.append(["", row["Typology"], row["Case Count"],
                row["Avg Amount (AUD)"], row["Total Amount (AUD)"], ""])
    r   = ws1.max_row
    alt = (i % 2 == 0)
    style_data_row(ws1, r, 6, alt)
    ws1.cell(r, 4).number_format = AUD
    ws1.cell(r, 5).number_format = AUD

set_col_widths(ws1, [3, 38, 22, 28, 22, 22])


# ─────────────────────────────────────────────
# SHEET 2: CONFIRMED LAUNDERING CASES
# ─────────────────────────────────────────────
ws2 = wb.create_sheet("Confirmed Cases")
ws2.sheet_view.showGridLines = False

ws2.merge_cells("A1:G1")
c = ws2["A1"]
c.value     = "Confirmed Money Laundering Cases — All 1,745 Transactions"
c.font      = TITLE_FONT
c.alignment = CENTER
c.fill      = PatternFill("solid", start_color="EBF3FB")
ws2.row_dimensions[1].height = 30
ws2.append([])

cols      = ["nameOrig", "nameDest", "amount", "type", "payment_method", "laundering_typology", "timestamp"]
headers   = ["Account (Origin)", "Account (Dest)", "Amount (AUD)", "Type", "Payment Method", "Typology", "Timestamp"]
col_widths = [18, 18, 18, 14, 18, 16, 26]

write_dataframe(
    ws2,
    confirmed[cols],
    start_row=3,
    col_headers=headers,
    col_widths=col_widths,
    amount_cols={3}
)


# ─────────────────────────────────────────────
# SHEET 3: ML FLAGGED (TOP 100)
# ─────────────────────────────────────────────
ws3 = wb.create_sheet("ML Flagged (Top 100)")
ws3.sheet_view.showGridLines = False

ws3.merge_cells("A1:F1")
c = ws3["A1"]
c.value     = "Top 100 ML-Flagged Transactions (by Risk Score)"
c.font      = TITLE_FONT
c.alignment = CENTER
c.fill      = PatternFill("solid", start_color="EBF3FB")
ws3.row_dimensions[1].height = 30
ws3.append([])

ml_cols    = ["nameOrig", "nameDest", "amount", "ml_risk_score", "laundering_typology", "timestamp"]
ml_headers = ["Account (Origin)", "Account (Dest)", "Amount (AUD)", "ML Risk Score", "Typology", "Timestamp"]
ml_widths  = [18, 18, 18, 16, 16, 26]

write_dataframe(
    ws3,
    ml_flagged[ml_cols],
    start_row=3,
    col_headers=ml_headers,
    col_widths=ml_widths,
    amount_cols={3}
)


# ─────────────────────────────────────────────
# SHEET 4: TTR BREACHES
# ─────────────────────────────────────────────
ws4 = wb.create_sheet("TTR Breaches")
ws4.sheet_view.showGridLines = False

ws4.merge_cells("A1:G1")
c = ws4["A1"]
c.value     = "Threshold Transaction Reports (TTR) — Transactions ≥ $10,000 AUD"
c.font      = TITLE_FONT
c.alignment = CENTER
c.fill      = PatternFill("solid", start_color="FFF2CC")
ws4.row_dimensions[1].height = 30

ws4.merge_cells("A2:G2")
note = ws4["A2"]
note.value     = "Under AML/CTF Act 2006, all cash transactions ≥ $10,000 AUD must be reported to AUSTRAC within 10 business days."
note.font      = Font(name="Arial", size=9, italic=True, color="7F6000")
note.alignment = LEFT
ws4.append([])

ttr_cols    = ["nameOrig", "nameDest", "amount", "type", "payment_method", "city", "timestamp"]
ttr_headers = ["Account (Origin)", "Account (Dest)", "Amount (AUD)", "Type", "Payment Method", "City", "Timestamp"]
ttr_widths  = [18, 18, 18, 14, 18, 14, 26]

write_dataframe(
    ws4,
    ttr[ttr_cols],
    start_row=4,
    col_headers=ttr_headers,
    col_widths=ttr_widths,
    amount_cols={3}
)


# ─────────────────────────────────────────────
# SHEET 5: STRUCTURING CANDIDATES
# ─────────────────────────────────────────────
ws5 = wb.create_sheet("Structuring Candidates")
ws5.sheet_view.showGridLines = False

ws5.merge_cells("A1:G1")
c = ws5["A1"]
c.value     = "Structuring Candidates — Transactions $9,000–$9,999 AUD"
c.font      = TITLE_FONT
c.alignment = CENTER
c.fill      = PatternFill("solid", start_color="FCE4D6")
ws5.row_dimensions[1].height = 30

ws5.merge_cells("A2:G2")
note = ws5["A2"]
note.value     = "Transactions deliberately kept below the $10,000 AUD TTR threshold may indicate structuring (smurfing) — a key AML/CTF red flag under AUSTRAC guidelines."
note.font      = Font(name="Arial", size=9, italic=True, color="843C0C")
note.alignment = LEFT
ws5.append([])

str_cols    = ["nameOrig", "nameDest", "amount", "type", "payment_method", "city", "timestamp"]
str_headers = ["Account (Origin)", "Account (Dest)", "Amount (AUD)", "Type", "Payment Method", "City", "Timestamp"]
str_widths  = [18, 18, 18, 14, 18, 14, 26]

write_dataframe(
    ws5,
    structuring[str_cols],
    start_row=4,
    col_headers=str_headers,
    col_widths=str_widths,
    amount_cols={3}
)

# ─────────────────────────────────────────────
# SAVE
# ─────────────────────────────────────────────
output_path = "AML_SAR_Report.xlsx"
wb.save(output_path)
print(f"SAR report saved to: {output_path}")
print(f"Sheets: {wb.sheetnames}")
